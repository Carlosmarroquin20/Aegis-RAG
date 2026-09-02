"""
Unit tests for the DOCX parser.

Test documents are generated in memory with python-docx, so no fixture files are
needed and the real extraction path is exercised end to end.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from typing import Any

import pytest
from docx import Document as DocxDocument

from aegis.domain.models.ingestion import SupportedMimeType
from aegis.domain.ports.document_parser import ParseError
from aegis.infrastructure.parsers.docx_parser import DocxParser


def _docx_bytes(build: Callable[[Any], None]) -> bytes:
    doc = DocxDocument()
    build(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class TestDocxParser:
    def test_supported_mime_types(self) -> None:
        assert SupportedMimeType.APPLICATION_DOCX in DocxParser().supported_mime_types()

    def test_flat_document_returns_single_section(self) -> None:
        def build(doc: Any) -> None:
            doc.add_paragraph("First paragraph.")
            doc.add_paragraph("Second paragraph.")

        docs = DocxParser().parse(_docx_bytes(build), "flat.docx")
        assert len(docs) == 1
        assert "First paragraph." in docs[0].content
        assert "Second paragraph." in docs[0].content
        assert docs[0].source == "flat.docx"

    def test_headings_split_into_sections(self) -> None:
        def build(doc: Any) -> None:
            doc.add_heading("Introduction", level=1)
            doc.add_paragraph("Intro body.")
            doc.add_heading("Details", level=1)
            doc.add_paragraph("Details body.")

        docs = DocxParser().parse(_docx_bytes(build), "structured.docx")
        assert len(docs) == 2
        assert docs[0].section == "Introduction"
        assert "Intro body." in docs[0].content
        assert docs[1].section == "Details"

    def test_corrupt_bytes_raise_parse_error(self) -> None:
        with pytest.raises(ParseError):
            DocxParser().parse(b"this is not a docx file", "bad.docx")

    def test_empty_document_raises_parse_error(self) -> None:
        def build(_doc: Any) -> None:
            return None  # no content

        with pytest.raises(ParseError):
            DocxParser().parse(_docx_bytes(build), "empty.docx")
